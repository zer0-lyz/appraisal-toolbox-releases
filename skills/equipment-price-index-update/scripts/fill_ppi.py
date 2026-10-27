# -*- coding: utf-8 -*-
"""
fill_ppi.py — 把 NBS 抓取结果回填进底稿「全年(更新此表数据）」

回填策略（保守，只补不覆盖）
    只填「底稿为空或为 0」而「NBS 有值」的格子。
    底稿已有值的一律不动——实测底稿的 D~K 列存 2 位小数（如 101.23），
    比 NBS 官网发布的 1 位小数（101.2）更精细，属于不同精度来源，不应覆盖。

用法
    python fill_ppi.py                      # 输出到 _out/<底稿名>_已补PPI.xlsx
    python fill_ppi.py --inplace            # 直接改原底稿（自动带时间戳备份）
    python fill_ppi.py --tolerance 0.1      # 把 ±0.1 内的差异也视为一致（默认 0.005）
"""
import os
import re
import sys
import json
import shutil
import warnings
from datetime import datetime

warnings.filterwarnings('ignore')
try:
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
except Exception:
    pass

import openpyxl

BASE = os.environ.get('FA_WORKSPACE') or os.getcwd()
OUT = os.path.join(BASE, '_out')
DRAFT = os.environ.get('FA_DRAFT') or os.path.join(
    BASE, '固定资产底稿模板', 'XXX公司-固定资产设备底稿-V1.0.xlsx')
SHEET = '全年(更新此表数据）'

COL2KEY = {
    'D': '家具制造业',
    'E': '文教、工美、体育和娱乐用品制造业',
    'F': '塑料制品',
    'G': '通用设备制造业',
    'H': '专用设备制造业',
    'I': '电气机械及器材制造业',
    'J': '通信设备、计算机及其他电子设备制造业',
    'K': '仪器仪表制造业',
    'L': '中国:PPI:汽车整车制造',
    'M': '中国:PPI:汽车车身、挂车制造',
}


def main():
    args = sys.argv[1:]
    inplace = '--inplace' in args
    draft = DRAFT
    js = os.path.join(OUT, 'nbs_ppi.json')
    output = ''
    tol = 0.005
    for i, a in enumerate(args):
        if a == '--draft' and i + 1 < len(args):
            draft = args[i + 1]
        if a == '--json' and i + 1 < len(args):
            js = args[i + 1]
        if a == '--output' and i + 1 < len(args):
            output = args[i + 1]
        if a == '--tolerance' and i + 1 < len(args):
            tol = float(args[i + 1])

    if not os.path.exists(js):
        print('缺少 NBS 数据 %s，请先运行 nbs_ppi.py' % js)
        return 3
    if not os.path.exists(draft):
        print('底稿不存在: %s' % draft)
        return 3

    with open(js, encoding='utf-8') as f:
        nbs = json.load(f)['series']

    if inplace:
        stamp = datetime.now().strftime('%Y%m%d-%H%M%S-%f')
        b, e = os.path.splitext(draft)
        backup = '%s-备份%s%s' % (b, stamp, e)
        shutil.copy2(draft, backup)
        target = draft
        print('已备份 -> %s' % os.path.basename(backup))
    else:
        if output:
            target = os.path.abspath(output)
            if os.path.normcase(target) == os.path.normcase(os.path.abspath(draft)):
                print('预览输出不能覆盖原底稿')
                return 4
            if os.path.exists(target):
                print('预览输出文件已存在，为保护已有结果未覆盖: %s' % target)
                return 4
            os.makedirs(os.path.dirname(target), exist_ok=True)
        else:
            os.makedirs(OUT, exist_ok=True)
            stem = os.path.splitext(os.path.basename(draft))[0]
            target = os.path.join(OUT, '%s_已补PPI.xlsx' % stem)
            if os.path.exists(target):
                print('预览输出文件已存在，为保护已有结果未覆盖: %s' % target)
                return 4
        shutil.copy2(draft, target)
        print('副本 -> %s' % os.path.basename(target))

    wb = openpyxl.load_workbook(target, keep_vba=target.lower().endswith('.xlsm'))
    ws = wb[SHEET]

    rows = {}
    for r in range(4, ws.max_row + 1):
        y, m = ws.cell(r, 2).value, ws.cell(r, 3).value
        if y is None or m is None:
            continue
        try:
            rows[(int(y), int(m))] = r
        except Exception:
            pass
    print('底稿 %d 个 (年,月) 行  %s ~ %s' % (len(rows), min(rows), max(rows)))
    print()

    filled = 0
    per_col = {}
    skipped_conflict = []
    for col, key in COL2KEY.items():
        ser = nbs.get(key) or {}
        n = 0
        for ym, val in sorted(ser.items()):
            y, m = int(ym[:4]), int(ym[4:6])
            r = rows.get((y, m))
            if r is None:
                continue
            cell = ws['%s%d' % (col, r)]
            cur = cell.value
            try:
                cur_n = float(cur) if cur is not None else None
            except Exception:
                cur_n = None
            if cur_n is None or cur_n == 0:
                try:
                    cell.value = float(val)
                    n += 1
                except Exception:
                    pass
            elif tol > 0:
                try:
                    if abs(cur_n - float(val)) > tol:
                        skipped_conflict.append((col, r, ym, cur_n, float(val)))
                except Exception:
                    pass
        if n:
            per_col[col] = n
            filled += n
            print('  %s %-30s 回填 %d 处' % (col, key[:28], n))

    print()
    print('合计回填 %d 处' % filled)
    if skipped_conflict:
        print('底稿已有值且与 NBS 差异超过容差(%s)的格子 %d 个（未改动，仅供核对）：'
              % (tol, len(skipped_conflict)))
        for col, r, ym, cv, nv in skipped_conflict[:10]:
            print('   %s%d  %s 底稿=%s NBS=%s' % (col, r, ym, cv, nv))

    wb.save(target)
    print('\n已保存 -> %s' % target)

    rep = {'target': target, 'filled': filled, 'per_col': per_col,
           'conflicts': len(skipped_conflict), 'tolerance': tol,
           'at': datetime.now().strftime('%Y-%m-%d %H:%M:%S')}
    with open(os.path.join(OUT, 'ppi_fill_report.json'), 'w', encoding='utf-8') as f:
        json.dump(rep, f, ensure_ascii=False, indent=2)
    return 0


if __name__ == '__main__':
    main()
