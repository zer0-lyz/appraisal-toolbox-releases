# -*- coding: utf-8 -*-
"""
ppi_baseline.py — PPI 数据基线（已抓取数据 + 底稿现有值的持久化快照）

为什么需要它
    全量抓取一次要请求 4 个时间分片 × 10 个指标，既慢又容易触发 NBS 限流。
    数据本身（2007 年至今的历史值）几乎不会变，没必要每次重抓。
    把"已抓取的全量 + 底稿现有值"固化成基线后：
      · 日常只需增量抓最近几期（校验历史有没有被 NBS 修正）
      · 再补底稿缺的那几期即可

基线内容
    nbs             已抓取的 NBS 全量序列（10 个指标 × 各期）
    draft_snapshot  底稿「全年(更新此表数据）」D~M 列的现值快照
    meta            最后一期、行数等关键指标

用法
    python ppi_baseline.py --build              # 从 nbs_ppi.json + 底稿建立基线
    python ppi_baseline.py --merge _out/nbs_ppi.json   # 把新抓数据合并进基线
    python ppi_baseline.py --status             # 查看基线状态

环境变量
    FA_BASELINE   基线文件路径（默认 <工作目录>/data/ppi_baseline.json）
"""
import os
import sys
import json
import shutil
import warnings
from datetime import datetime

warnings.filterwarnings('ignore')
try:
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
except Exception:
    pass

import openpyxl

BASE = os.environ.get('FA_WORKSPACE') or os.getcwd()
BASELINE = os.environ.get('FA_BASELINE') or os.path.join(BASE, 'data', 'ppi_baseline.json')
DRAFT = os.environ.get('FA_DRAFT') or os.path.join(
    BASE, '固定资产底稿模板', 'XXX公司-固定资产设备底稿-V1.0.xlsx')
SHEET = '全年(更新此表数据）'
FIRST_DATA_ROW = 4

COL2KEY = {
    'D': '家具制造业',
    'E': '文教、工美、体育和娱乐用品制造业',
    'F': '塑料制品',
    'G': '通用设备制造业',
    'H': '专用设备制造业',
    'I': '电气机械及器材制造业',
    'J': '通信设备、计算机及其他电子设备制造业',
    'K': '仪器仪表制造业',
    'L': '中国:PPI:汽车整车制造',
    'M': '中国:PPI:汽车车身、挂车制造',
}
KEY2COL = {v: k for k, v in COL2KEY.items()}


def read_draft_snapshot(draft):
    """读出底稿 D~M 列的现值，返回 ({列: {yyyymm: 值}}, meta)"""
    if not os.path.exists(draft):
        return {}, {}
    wb = openpyxl.load_workbook(draft, data_only=True)
    ws = wb[SHEET]
    snap = {c: {} for c in COL2KEY}
    rows = 0
    last_period = None
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        y, m = ws.cell(r, 2).value, ws.cell(r, 3).value
        if y is None or m is None:
            continue
        try:
            ym = '%04d%02d' % (int(y), int(m))
        except Exception:
            continue
        rows += 1
        last_period = ym
        for col in COL2KEY:
            v = ws.cell(r, openpyxl.utils.column_index_from_string(col)).value
            if isinstance(v, (int, float)) and v != 0:
                snap[col][ym] = v
    return snap, {'draft_rows': rows, 'draft_last_period': last_period,
                  'draft_path': draft}


def load(p):
    if not os.path.exists(p):
        return None
    with open(p, encoding='utf-8') as f:
        return json.load(f)


def save(obj, p):
    os.makedirs(os.path.dirname(p), exist_ok=True)
    with open(p, 'w', encoding='utf-8') as f:
        json.dump(obj, f, ensure_ascii=False, separators=(',', ':'))
    return os.path.getsize(p)


def find_seed():
    """查找技能自带的初始基线（references/ppi_baseline.json，随技能包分发）。

    用途：让新环境**开箱即用**。即使只抓到最近几期，也能把 2007 年以来的
    历史序列补齐，不必先全量抓一次。
    返回文件路径；找不到返回 None（照常走全量抓取）。
    """
    here = os.path.dirname(os.path.abspath(__file__))
    cands = []
    ev = os.environ.get('FA_BASELINE_SEED')
    if ev:
        cands.append(ev)
    # 技能目录内运行：scripts/../references/ppi_baseline.json
    cands.append(os.path.join(here, os.pardir, 'references', 'ppi_baseline.json'))
    # 脚本被复制到项目 tools/ 时的同目录副本
    cands.append(os.path.join(here, 'ppi_baseline.seed.json'))
    for c in cands:
        c = os.path.abspath(c)
        if os.path.exists(c) and c != os.path.abspath(BASELINE):
            return c
    return None


def build(files):
    nbs_all = {}
    src = None
    # 先用技能自带的初始基线打底（若有）——保证历史完整，新环境开箱即用
    seed = find_seed()
    if seed:
        sd = load(seed)
        if sd and sd.get('nbs'):
            for k, v in sd['nbs'].items():
                nbs_all.setdefault(k, {}).update(v)
            src = sd.get('source')
            print('打底：技能自带基线 %s' % os.path.basename(seed))
            print('      %d 个序列 / %d 个数据点，最新期 %s'
                  % (len(sd['nbs']), sum(len(x) for x in sd['nbs'].values()),
                     (sd.get('meta') or {}).get('nbs_last_period')))
    for f in files:
        d = load(f)
        if not d:
            continue
        if not src:
            src = d.get('source')
        for k, v in (d.get('series') or {}).items():
            nbs_all.setdefault(k, {}).update(v)     # 后抓的覆盖先抓的
    if not nbs_all:
        print('没有可用的 NBS 数据')
        return 1
    draft_snap, meta = read_draft_snapshot(DRAFT)
    nbs_last = max((max(v) for v in nbs_all.values() if v), default=None)
    obj = {
        'schema_version': 1,
        'generated_at': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
        'source': src or '国家统计局 分行业工业生产者出厂价格指数(上年同月=100)',
        'note': '已抓取的 NBS 全量数据 + 底稿现有值快照。日常增量抓取只需请求最近几期做校验。',
        'nbs': {k: dict(sorted(v.items())) for k, v in sorted(nbs_all.items())},
        'draft_snapshot': {k: dict(sorted(v.items())) for k, v in sorted(draft_snap.items())},
        'meta': {
            'nbs_last_period': nbs_last,
            'nbs_series_count': len(nbs_all),
            'nbs_total_points': sum(len(v) for v in nbs_all.values()),
            'draft_rows': meta.get('draft_rows'),
            'draft_last_period': meta.get('draft_last_period'),
            'column_map': COL2KEY,
        },
    }
    sz = save(obj, BASELINE)
    print('基线已建立 -> %s (%.1f KB)' % (BASELINE, sz / 1024.0))
    print()
    print('NBS 序列 (%d 个, 共 %d 个数据点):' % (len(nbs_all), obj['meta']['nbs_total_points']))
    for k, v in obj['nbs'].items():
        ks = sorted(v)
        print('   %-34s %4d 期  %s ~ %s' % (k[:32], len(v), ks[0], ks[-1]))
    print()
    if meta.get('draft_rows'):
        print('底稿快照: %d 行, 最后一期 %s'
              % (meta['draft_rows'], meta.get('draft_last_period') or '-'))
        for col in COL2KEY:
            s = obj['draft_snapshot'].get(col) or {}
            if s:
                ks = sorted(s)
                print('   %s %-30s %4d 期  %s ~ %s'
                      % (col, COL2KEY[col][:28], len(s), ks[0], ks[-1]))
    else:
        print('底稿快照: 未读取（底稿文件不存在或路径不对——不影响 NBS 数据）')
    return 0


def merge(files):
    obj = load(BASELINE)
    if not obj:
        print('基线不存在，请先 --build')
        return 1
    # 先备份基线（合并是就地改，出问题可回退）
    stamp = datetime.now().strftime('%Y%m%d-%H%M%S-%f')
    root, ext = os.path.splitext(BASELINE)
    b = '%s-备份%s%s' % (root, stamp, ext)
    shutil.copy2(BASELINE, b)
    added = 0
    changed = 0
    for f in files:
        d = load(f)
        if not d:
            continue
        for k, v in (d.get('series') or {}).items():
            tgt = obj['nbs'].setdefault(k, {})
            for ym, val in v.items():
                if ym not in tgt:
                    tgt[ym] = val
                    added += 1
                elif abs(float(tgt[ym]) - float(val)) > 0.0001:
                    tgt[ym] = val      # NBS 修正了历史值
                    changed += 1
    for k in obj['nbs']:
        obj['nbs'][k] = dict(sorted(obj['nbs'][k].items()))
    obj['generated_at'] = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    obj['meta']['nbs_last_period'] = max(
        (max(v) for v in obj['nbs'].values() if v), default=None)
    obj['meta']['nbs_total_points'] = sum(len(v) for v in obj['nbs'].values())
    save(obj, BASELINE)
    print('基线已合并 -> %s (备份 %s)' % (BASELINE, os.path.basename(b)))
    print('   新增 %d 期，NBS 修正历史值 %d 期' % (added, changed))
    print('   最新期: %s' % obj['meta']['nbs_last_period'])
    return 0


def status():
    seed = find_seed()
    obj = load(BASELINE)
    if not obj:
        print('基线不存在: %s' % BASELINE)
        if seed:
            print('   但找到技能自带基线: %s' % seed)
            print('   可直接执行 --build 用它打底（无需全量抓取，几秒钟完成）')
        else:
            print('   且未找到技能自带基线，请先运行 nbs_ppi.py 全量抓取')
        return 1
    m = obj['meta']
    print('基线文件: %s' % BASELINE)
    print('   建立时间      : %s' % obj['generated_at'])
    print('   技能自带基线  : %s' % (os.path.basename(seed) if seed else '（未找到）'))
    print('   来源          : %s' % str(obj['source'])[:66])
    print('   NBS 序列      : %d 个, %d 个数据点, 最新期 %s'
          % (m['nbs_series_count'], m['nbs_total_points'], m['nbs_last_period']))
    print('   底稿快照      : %s 行, 最后一期 %s'
          % (m.get('draft_rows') or '未读取', m.get('draft_last_period') or '-'))
    print()
    print('   %-34s %-22s %s' % ('序列', 'NBS 期范围', '底稿期范围'))
    print('   ' + '-' * 74)
    for k, col in KEY2COL.items():
        nv = obj['nbs'].get(k) or {}
        dv = obj['draft_snapshot'].get(col) or {}
        nr = '%s~%s' % (min(nv), max(nv)) if nv else '-'
        dr = '%s~%s' % (min(dv), max(dv)) if dv else '-'
        print('   %-34s %-22s %s' % (k[:32], nr, dr))
    return 0


def main():
    args = sys.argv[1:]
    files = [a for a in args if a.endswith('.json')]
    if '--build' in args:
        if not files:
            files = [os.path.join(BASE, '_out', 'nbs_ppi.json')]
        return build(files)
    if '--merge' in args:
        if not files:
            files = [os.path.join(BASE, '_out', 'nbs_ppi.json')]
        return merge(files)
    if '--status' in args:
        return status()
    print(__doc__)
    return 0


if __name__ == '__main__':
    sys.exit(main())
