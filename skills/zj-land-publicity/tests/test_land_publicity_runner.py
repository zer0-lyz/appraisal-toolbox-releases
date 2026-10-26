import os
import json
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl
from unittest.mock import patch

sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))

from land_publicity_runner import (
    OUTPUT_COLUMNS, _land_use_categories, _land_use_matches_request, _list_record_matches_request,
    _history_item_with_price_notice, _list_server_filters, _plot_ratio_text, _website_server_filters,
    enrich_record, execute_request, filter_records, load_history_file, row_from_item, validate_request, write_excel, write_result_html,
)
from land_price import normalize_land_price, normalize_price_text, price_value, resolve_detail_price
from scrape_zj_land import build_map_assets, fetch_records, fetch_region_codes, parse_content_fields, save_to_excel


class JSONResponseFixture:
    status_code = 200
    headers = {"Content-Type": "application/json"}

    def raise_for_status(self):
        pass

    @property
    def text(self):
        payload = {"code": 0, **self.json()}
        if isinstance(payload.get("data"), dict) and "records" in payload["data"]:
            payload["data"].setdefault("total", 1000)
        return json.dumps(payload, ensure_ascii=False)


class LandPublicityDateTests(unittest.TestCase):
    def test_server_filters_use_region_tree_and_publish_window(self):
        request = {"district": "宁波市", "location": "镇海区", "provinceWide": False, "startDate": "2025-01-01", "endDate": "2025-01-31"}
        filters = _list_server_filters(request)
        self.assertEqual(filters["regionName"], "镇海区")
        self.assertEqual(filters["fallbackRegionName"], "宁波市")
        self.assertEqual(filters["enrollStartTime"], 1735660800000)
        self.assertEqual(filters["nowTime"], 1738339199999)

    def test_region_tree_resolves_all_descendant_codes(self):
        class Response(JSONResponseFixture):
            def json(self):
                return {"data": [{"districtName": "浙江省", "districtCode": "330000", "children": [
                    {"districtName": "宁波市", "districtCode": "330200", "children": [
                        {"districtName": "镇海区", "districtCode": "330211", "children": []},
                    ]},
                ]}]}

        with patch("scrape_zj_land.requests.get", return_value=Response()):
            self.assertEqual(fetch_region_codes("宁波市"), ["330200", "330211"])

    def test_server_filters_are_sent_to_publicity_list(self):
        calls = []

        class Response(JSONResponseFixture):
            def json(self):
                return {"data": {"records": []}}

        with patch("scrape_zj_land.requests.get", side_effect=lambda url, **kwargs: (calls.append(url) or Response())):
            fetch_records(
                max_pages=1,
                server_filters={
                    "regionCode": "330200,330211",
                    "publishStartTime": 1735660800000,
                    "publishEndTime": 1738339199999,
                },
            )

        self.assertEqual(len(calls), 1)
        self.assertIn("regionCode=330200%2C330211", calls[0])
        self.assertIn("publishStartTime=1735660800000", calls[0])
        self.assertIn("publishEndTime=1738339199999", calls[0])

    def test_county_region_takes_precedence_over_city_fallback(self):
        calls = []

        class Response(JSONResponseFixture):
            def json(self):
                return {"data": {"records": []}}

        with patch("scrape_zj_land.fetch_region_codes", return_value=["330211"]), patch(
            "scrape_zj_land.requests.get", side_effect=lambda url, **kwargs: (calls.append(url) or Response())
        ):
            fetch_records(max_pages=1, server_filters={"regionName": "镇海区", "fallbackRegionName": "宁波市"})

        self.assertIn("regionCode=330211", calls[0])

    def test_official_land_use_codes_cover_all_four_categories(self):
        fixtures = [
            ("住宅用地", {"NAME_": "普通商品住房用地（二类）", "KEY_": "V3-07010203", "PARENTKEY_": "V3-07"}),
            ("商服用地", {"NAME_": "零售商业用地", "KEY_": "V3-090101", "PARENTKEY_": "V3-09"}),
            ("工矿仓储", {"NAME_": "一类工业用地（其他工业用地）", "KEY_": "V3-10010102", "PARENTKEY_": "V3-10"}),
            ("其他用地", {"NAME_": "社会停车场用地", "KEY_": "V3-120803", "PARENTKEY_": "V3-12"}),
            ("其他用地", {"NAME_": "供电用地", "KEY_": "V3-1303", "PARENTKEY_": "V3-13"}),
            ("其他用地", {"NAME_": "公园绿地", "KEY_": "V3-1401", "PARENTKEY_": "V3-14"}),
            ("住宅用地", {"NAME_": "普通商品房", "KEY_": "07010201"}),
            ("商服用地", {"NAME_": "零售商业用地", "KEY_": "0501"}),
            ("工矿仓储", {"NAME_": "其他工业用地", "KEY_": "060102"}),
            ("其他用地", {"NAME_": "教育用地", "KEY_": "0803"}),
        ]
        categories = ["住宅用地", "商服用地", "工矿仓储", "其他用地"]
        for expected, plan_use in fixtures:
            record = {"_sourceEndpoint": "landbidding", "planUse": json.dumps([plan_use], ensure_ascii=False)}
            self.assertEqual(_land_use_categories(record, {}), {expected})
            other = next(value for value in categories if value != expected)
            self.assertTrue(_land_use_matches_request({"landUses": [expected, other]}, record, {}, allow_unknown=False))
            choices_without_match = [value for value in categories if value != expected][:2]
            self.assertFalse(_land_use_matches_request({"landUses": choices_without_match}, record, {}, allow_unknown=False))

    def test_single_official_land_use_filter_is_used_for_each_category(self):
        codes = {"住宅用地": 1, "商服用地": 2, "工矿仓储": 3, "其他用地": 4}
        for land_use, code in codes.items():
            request = {"landUses": [land_use]}
            self.assertEqual(_list_server_filters(request).get("landUse"), None)
            self.assertEqual(_website_server_filters(request)["landUse"], code)

    def test_plot_ratio_is_preserved_in_excel_and_html(self):
        raw_ratio = '{"RJL_X_FH":"≤","RJL_X":"1.01","RJL_S_FH":"≤","RJL_S":"1.6"}'
        self.assertEqual(_plot_ratio_text(raw_ratio), "X≤1.01；S≤1.6")
        enriched = enrich_record({"sourceId": "ratio-fixture", "_sourceEndpoint": "landbidding", "rjl": raw_ratio}, lambda _source_id: {"plotRatio": "{\"RJL_X_FH\":\"≤\",\"RJL_X\":\"9\"}"})
        self.assertEqual(enriched["plot_ratio"], "X≤1.01；S≤1.6")
        detail_fallback = enrich_record({"sourceId": "ratio-fixture", "_sourceEndpoint": "landbidding"}, lambda _source_id: {"plotRatio": raw_ratio})
        self.assertEqual(detail_fallback["plot_ratio"], "X≤1.01；S≤1.6")
        item = {
            "record": {"publicityId": "2026-29", "sourceCode": "2026-29", "rjl": raw_ratio},
            "detail": {}, "detail_data": {"plotRatio": raw_ratio}, "coord": {},
            "location": "余姚市", "land_use": "普通商品住房用地（二类）", "area_mu": None,
            "area_sqm": None, "start_price": None, "start_total_price": "", "deal_price": None,
            "deal_total_price": "", "detail_url": "https://example.invalid/detail", "transferee": "",
            "trade_form": "", "term": "", "plot_ratio": "",
        }
        row = row_from_item(item)
        self.assertEqual(row["容积率"], "X≤1.01；S≤1.6")
        self.assertIn("容积率", OUTPUT_COLUMNS)
        with tempfile.TemporaryDirectory() as directory:
            excel_path = os.path.join(directory, "result.xlsx")
            html_path = os.path.join(directory, "result.html")
            write_excel(Path(excel_path), [row], {"maxPages": 1}, {})
            workbook = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
            values = list(workbook.active.iter_rows(values_only=True))
            workbook.close()
            ratio_column = values[0].index("容积率")
            self.assertEqual(values[1][ratio_column], "X≤1.01；S≤1.6")
            write_result_html(Path(html_path), [row], {"maxPages": 1}, {}, {"map": "", "excel": "result.xlsx"})
            html_text = Path(html_path).read_text(encoding="utf-8")
            self.assertIn("容积率", html_text)
            self.assertIn("X≤1.01；S≤1.6", html_text)

    def test_plot_ratio_is_preserved_in_map_popup(self):
        with tempfile.TemporaryDirectory() as directory:
            output_path = Path(directory) / "result.xlsx"
            rows = [{
                "sourceCode": "2026-26", "districtName": "余姚市", "releaseTime": "2026-10-22",
                "_detail_url": "https://example.invalid/detail", "_location": "余姚市",
                "_use": "零售商业用地", "_plot_ratio": "X≤1.3；S≤1.4",
                "_coord": {"center": {"lng": 121.0, "lat": 30.0}},
            }]
            with patch("scrape_zj_land.load_map_config", return_value={"amapEnabled": False}), patch(
                "scrape_zj_land.browser_script", return_value=""
            ):
                assets = build_map_assets(rows, output_path)
            points = json.loads(Path(assets["coord_json"]).read_text(encoding="utf-8"))
            self.assertEqual(points[0]["items"][0]["plotRatio"], "X≤1.3；S≤1.4")
            map_html = Path(assets["map_html"]).read_text(encoding="utf-8")
            self.assertIn("容积率", map_html)
            self.assertIn("row.plotRatio", map_html)

    def test_list_candidate_filter_uses_list_content_before_detail_fetch(self):
        request = {
            "landUses": ["工矿仓储"],
            "startDate": "2025-01-01",
            "endDate": "2025-12-31",
        }
        matching = {"releaseTime": "2025年06月01日", "content": "<td>土地用途</td><td>工矿仓储用地</td>"}
        non_matching = {"releaseTime": "2025年06月01日", "content": "<td>土地用途</td><td>住宅用地</td>"}
        self.assertTrue(_list_record_matches_request(request, matching))
        self.assertFalse(_list_record_matches_request(request, non_matching))

    def test_list_fetch_filters_candidates_and_stops_at_start_date(self):
        pages = [
            [
                {"id": "keep", "releaseTime": "2025-01-02"},
                {"id": "drop", "releaseTime": "2025-01-01"},
            ],
            [{"id": "too-old", "releaseTime": "2024-12-31"}],
        ]
        calls = []

        class Response(JSONResponseFixture):
            def __init__(self, records):
                self.records = records

            def json(self):
                return {"data": {"records": self.records}}

        def get(url, **kwargs):
            calls.append(url)
            return Response(pages[len(calls) - 1])

        with patch("scrape_zj_land.requests.get", side_effect=get), patch("scrape_zj_land.time.sleep"):
            rows = fetch_records(
                max_pages=10,
                record_filter=lambda record: record["releaseTime"] >= "2025-01-01",
                stop_before="2025-01-01",
            )

        self.assertEqual([row["id"] for row in rows], ["keep", "drop"])
        self.assertEqual(len(calls), 2)
        self.assertIn("size=500", calls[0])

    def test_release_date_uses_inclusive_range(self):
        request = {
            "tradeForm": "",
            "tradeMethods": [],
            "tradeStages": [],
            "district": "",
            "provinceWide": True,
            "location": "",
            "landUses": [],
            "startDate": "2025-01-01",
            "endDate": "2025-01-31",
            "startYear": "",
            "quotePreset": "all",
            "quoteStartDate": "",
            "quoteEndDate": "",
            "startPriceMin": None,
            "startPriceMax": None,
            "areaMin": None,
            "areaMax": None,
            "areaUnit": "sqm",
        }
        records = [
            {"releaseTime": "2024-12-31"},
            {"releaseTime": "2025-01-01"},
            {"releaseTime": "2025-01-31"},
            {"releaseTime": "2025-02-01"},
        ]
        enriched = [
            {"record": record, "detail": {}, "location": "", "start_price": None, "area_sqm": None, "area_mu": None}
            for record in records
        ]
        rows, _ = filter_records(enriched, request)
        self.assertEqual([row["record"]["releaseTime"] for row in rows], ["2025-01-01", "2025-01-31"])
        validated = validate_request({"outputDirectory": os.getcwd(), "provinceWide": True, **request})
        self.assertEqual(validated["maxPages"], 200)

    def test_start_date_cannot_follow_end_date(self):
        with self.assertRaisesRegex(ValueError, "LAND_DATE_RANGE_INVALID"):
            validate_request({
                "outputDirectory": os.getcwd(),
                "provinceWide": True,
                "startDate": "2025-02-01",
                "endDate": "2025-01-01",
            })


class LandPriceNormalizationTests(unittest.TestCase):
    def test_official_total_and_unit_price_examples(self):
        total_first = normalize_land_price(219.35, "万元", 1948, source="official-fixture")
        self.assertEqual(total_first["verificationStatus"], "verified")
        self.assertEqual(total_first["totalPriceWanDisplay"], "219.35")
        self.assertEqual(total_first["unitPriceYuanSqmDisplay"], "1126.03")

        unit_first = normalize_land_price(900, "元/平方米", 8059, source="official-fixture")
        self.assertEqual(unit_first["unitPriceYuanSqmDisplay"], "900")
        self.assertEqual(unit_first["totalPriceWanDisplay"], "725.31")

    def test_missing_unknown_and_conflicting_units_are_not_guessed(self):
        missing = normalize_land_price(900, None, 8059)
        unknown = normalize_land_price(900, "元/亩", 8059)
        conflict = normalize_land_price("900元/平方米", "万元", 8059)
        unsupported_inline_conflict = normalize_land_price("900万元/亩", "万元", 8059)
        multiple = normalize_price_text("成交价900元/平方米；起始价1,000万元", 1000)

        self.assertEqual(missing["verificationStatus"], "missing_unit")
        self.assertEqual(unknown["verificationStatus"], "unknown_unit")
        self.assertEqual(conflict["verificationStatus"], "unit_conflict")
        self.assertEqual(unsupported_inline_conflict["verificationStatus"], "unit_conflict")
        self.assertIsNone(unsupported_inline_conflict["totalPriceWanDisplay"])
        self.assertEqual(multiple["verificationStatus"], "source_conflict")
        self.assertIsNone(multiple["unitPriceYuanSqmDisplay"])
        self.assertIsNone(multiple["totalPriceWanDisplay"])

    def test_matching_deal_price_list_keeps_other_quotes_as_evidence(self):
        detail = {
            "dealPrice": 900,
            "dealPriceList": [
                {"price": 800, "unit": "元/平方米"},
                {"price": 900, "unit": "元/平方米"},
            ],
        }
        resolved = resolve_detail_price(detail, "deal", 1000)
        self.assertEqual(resolved["verificationStatus"], "verified")
        self.assertEqual(price_value(resolved, "unit"), 900.0)
        self.assertEqual(len(resolved["evidence"]), 2)
        self.assertIn("唯一匹配", resolved["selectionBasis"])

        ambiguous = resolve_detail_price({
            "dealPrice": 900,
            "dealPriceList": [
                {"price": 900, "unit": "元/平方米"},
                {"price": 900, "unit": "万元"},
            ],
        }, "deal", 1000)
        self.assertEqual(ambiguous["verificationStatus"], "source_conflict")

    def test_zero_price_is_preserved_and_missing_area_keeps_explicit_total(self):
        zero = normalize_land_price(0, "元/平方米", 0)
        total_without_area = normalize_land_price(219.35, "万元", None)
        self.assertEqual(price_value(zero, "unit"), 0.0)
        self.assertEqual(price_value(zero, "total"), 0.0)
        self.assertEqual(total_without_area["verificationStatus"], "verified_partial")
        self.assertEqual(price_value(total_without_area, "total"), 219.35)
        self.assertEqual(price_value(total_without_area, "unit"), "")

    def test_start_price_filter_excludes_unverified_price_and_keeps_zero(self):
        request = validate_request({
            "outputDirectory": os.getcwd(),
            "provinceWide": True,
            "startPriceMin": 0,
            "startPriceMax": 1200,
        })
        missing_unit = enrich_record(
            {"startPrice": 1000, "landArea": 20000},
            lambda _source_id: {},
        )
        filtered, summary = filter_records([missing_unit], request)
        self.assertEqual(filtered, [])
        self.assertEqual(summary["unknown"]["unknownStartPrice"], 1)

        zero_price = enrich_record(
            {"startPrice": 0, "startPriceUnit": "元/平方米", "landArea": 20000},
            lambda _source_id: {},
        )
        request["startPriceMin"] = 0
        request["startPriceMax"] = 0
        filtered, summary = filter_records([zero_price], request)
        self.assertEqual(len(filtered), 1)
        self.assertEqual(filtered[0]["start_price"], 0)
        self.assertEqual(summary["unknown"]["unknownStartPrice"], 0)

    def test_old_history_without_price_metadata_keeps_numbers_and_marks_them(self):
        legacy_item = {
            "record": {"publicityId": "old-1", "sourceCode": "old-1", "districtName": "余姚市"},
            "detail": {}, "detail_data": {}, "coord": {}, "source_id": "old-1",
            "detail_url": "", "location": "旧地块", "land_use": "工业用地",
            "area_mu": None, "area_sqm": None,
            "start_price": 12.34, "start_total_price": 56.78,
            "deal_price": 90.12, "deal_total_price": 34.56,
        }
        with tempfile.TemporaryDirectory() as directory:
            history_path = Path(directory) / "old-history.json"
            history_path.write_text(json.dumps({
                "type": "zj-land-publicity-history", "version": 1, "items": [legacy_item],
            }, ensure_ascii=False), encoding="utf-8")
            loaded = load_history_file(history_path)

        prepared = _history_item_with_price_notice(loaded["items"][0])
        row = row_from_item(prepared)
        self.assertEqual(row["起始单价(元/平方米)"], 12.34)
        self.assertEqual(row["起始总价(万元)"], 56.78)
        self.assertEqual(row["成交单价(元/平方米)"], 90.12)
        self.assertEqual(row["成交总价(万元)"], 34.56)
        self.assertEqual(row["起始价核验状态"], "旧历史未核验")
        self.assertEqual(row["成交价核验状态"], "旧历史未核验")

        partial = _history_item_with_price_notice({**legacy_item, "deal_price_meta": {}})
        self.assertEqual(partial["deal_price_meta"]["verificationStatus"], "legacy_unverified")

        with tempfile.TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "old-history.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.append([
                "公示编号", "公示标题", "行政区", "土地面积(亩)", "土地面积(平方米)",
                "起始单价(元/平方米)", "起始总价(万元)", "成交单价(元/平方米)",
                "成交总价(万元)", "详情页网址",
            ])
            sheet.append(["xlsx-old", "xlsx-old", "宁波市", 2.3, 1500, 8.5, 1.2, 9.5, 1.3, ""])
            workbook.save(workbook_path)
            workbook.close()
            loaded_xlsx = load_history_file(workbook_path)

        xlsx_item = _history_item_with_price_notice(loaded_xlsx["items"][0])
        xlsx_row = row_from_item(xlsx_item)
        self.assertEqual(xlsx_row["起始单价(元/平方米)"], 8.5)
        self.assertEqual(xlsx_row["起始总价(万元)"], 1.2)
        self.assertEqual(xlsx_row["成交单价(元/平方米)"], 9.5)
        self.assertEqual(xlsx_row["成交总价(万元)"], 1.3)
        self.assertEqual(xlsx_row["成交价核验状态"], "旧历史未核验")

    def test_reused_old_history_summary_explains_refresh_action(self):
        old_item = {
            "record": {"publicityId": "old-reuse", "sourceCode": "old-reuse", "districtName": "宁波市"},
            "detail": {}, "detail_data": {}, "coord": {}, "source_id": "old-reuse",
            "detail_url": "", "location": "旧地块", "land_use": "工业用地",
            "area_mu": None, "area_sqm": 1000,
            "start_price": 12.0, "start_total_price": 1.2,
            "deal_price": 15.0, "deal_total_price": 1.5,
        }
        with tempfile.TemporaryDirectory() as directory:
            history_path = Path(directory) / "legacy.json"
            history_path.write_text(json.dumps({
                "type": "zj-land-publicity-history", "version": 1, "items": [old_item],
            }, ensure_ascii=False), encoding="utf-8")
            result = execute_request({
                "outputDirectory": directory,
                "historyPath": str(history_path),
                "historyRefresh": False,
                "provinceWide": True,
                "generateMap": False,
                "maxPages": 1,
            }, fetcher=lambda **_kwargs: self.fail("history reuse must not fetch the list"),
               detail_fetcher=lambda _source_id: self.fail("history reuse must not fetch detail"),
               progress=lambda *_args, **_kwargs: None)

        warnings = result["filterSummary"]["warnings"]
        self.assertTrue(any("旧历史价格单位未核验" in warning and "重新抓取详情复核" in warning for warning in warnings))


class LegacyLandExportTests(unittest.TestCase):
    def test_explicit_html_area_and_price_units_feed_runner_and_legacy_export(self):
        content = (
            "<td>土地面积（平方米）</td><td>20000</td>"
            "<td>成交价（万元）</td><td>219.35</td>"
        )
        detail = parse_content_fields(content)
        self.assertEqual(detail["土地面积(平方米)"], "20000")
        self.assertEqual(detail["土地面积(亩)"], "")
        self.assertEqual(detail["成交价"], "219.35")
        self.assertEqual(detail["成交价单位"], "万元")

        enriched = enrich_record({"content": content}, lambda _source_id: {})
        self.assertEqual(enriched["area_sqm"], 20000.0)
        self.assertIsNone(enriched["area_mu"])
        self.assertEqual(enriched["deal_total_price"], 219.35)
        self.assertEqual(enriched["deal_price"], 109.68)

        mu_detail = parse_content_fields("<td>土地面积(亩)</td><td>3.5</td>")
        unknown_detail = parse_content_fields("<td>土地面积（公顷）</td><td>1</td>")
        unsupported_compounds = [
            parse_content_fields("<td>土地面积（公亩）</td><td>1</td>"),
            parse_content_fields("<td>土地面积（万平方米）</td><td>1</td>"),
        ]
        unknown_price_unit = parse_content_fields("<td>成交价（万元/亩）</td><td>100</td>")
        self.assertEqual(mu_detail["土地面积(亩)"], "3.5")
        self.assertEqual(unknown_detail["土地面积(亩)"], "")
        self.assertEqual(unknown_detail["土地面积(平方米)"], "")
        self.assertEqual(unknown_detail["土地面积原始单位"], "公顷")
        for parsed_area in unsupported_compounds:
            self.assertEqual(parsed_area["土地面积(亩)"], "")
            self.assertEqual(parsed_area["土地面积(平方米)"], "")
        self.assertEqual(unknown_price_unit["成交价单位"], "万元/亩")
        unknown_price = enrich_record(
            {"content": "<td>成交价（万元/亩）</td><td>100</td>"},
            lambda _source_id: {},
        )
        self.assertEqual(unknown_price["deal_price_meta"]["verificationStatus"], "unknown_unit")
        self.assertEqual(unknown_price["deal_total_price"], "")
        self.assertEqual(unknown_price["deal_price"], "")

        records = [
            {"publicityId": "html-1", "sourceCode": "html-1", "districtName": "余姚市", "content": content},
            {
                "_sourceEndpoint": "landbidding",
                "publicityId": "zero-1", "sourceCode": "zero-1", "districtName": "余姚市",
                "landArea": 1000, "cjj": 0, "cjfs": "元/平方米",
                "content": "<td>起始价（万元）</td><td>=1+1</td>",
            },
            {
                "publicityId": "unknown-area", "sourceCode": "unknown-area", "districtName": "余姚市",
                "content": "<td>土地面积（公顷）</td><td>1</td><td>成交价（万元）</td><td>100</td>",
            },
            {
                "publicityId": "unknown-price-unit", "sourceCode": "unknown-price-unit", "districtName": "余姚市",
                "content": "<td>成交价（万元/亩）</td><td>100</td>",
            },
        ]
        with tempfile.TemporaryDirectory() as directory:
            output_path = Path(directory) / "legacy-export.xlsx"
            with patch("builtins.print"):
                save_to_excel(records, output_path)
            workbook = openpyxl.load_workbook(output_path, data_only=False)
            sheet = workbook.active
            headers = [cell.value for cell in sheet[1]]
            column = {name: headers.index(name) + 1 for name in headers}

            self.assertEqual(sheet.cell(2, column["土地面积(平方米)"]).value, 20000.0)
            self.assertIn(sheet.cell(2, column["土地面积(亩)"]).value, (None, ""))
            self.assertEqual(sheet.cell(2, column["成交单价(元/平方米)"]).value, 109.68)
            self.assertEqual(sheet.cell(2, column["成交总价(万元)"]).value, 219.35)
            self.assertEqual(sheet.cell(2, column["成交价原始单位"]).value, "万元")
            self.assertEqual(sheet.cell(2, column["成交价核验状态"]).value, "已核验")
            self.assertEqual(sheet.cell(3, column["成交单价(元/平方米)"]).value, 0)
            self.assertEqual(sheet.cell(3, column["成交总价(万元)"]).value, 0)
            self.assertEqual(sheet.cell(3, column["起始价原始金额"]).value, "=1+1")
            self.assertEqual(sheet.cell(3, column["起始价原始金额"]).data_type, "s")
            self.assertIn(sheet.cell(4, column["土地面积(平方米)"]).value, (None, ""))
            self.assertIn(sheet.cell(4, column["土地面积(亩)"]).value, (None, ""))
            self.assertEqual(sheet.cell(4, column["成交总价(万元)"]).value, 100)
            self.assertIn("面积单位未核验", sheet.cell(4, column["成交价换算依据"]).value)
            self.assertEqual(sheet.cell(5, column["成交价原始单位"]).value, "万元/亩")
            self.assertEqual(sheet.cell(5, column["成交价核验状态"]).value, "单位未知")
            self.assertIn(sheet.cell(5, column["成交总价(万元)"]).value, (None, ""))
            workbook.close()

    def test_legacy_map_keeps_price_provenance(self):
        content = "<td>土地面积（平方米）</td><td>20000</td><td>成交价（万元）</td><td>219.35</td>"
        record = {"publicityId": "map-html-1", "sourceCode": "map-html-1", "content": content, "sourceId": "map-id"}
        with tempfile.TemporaryDirectory() as directory:
            output_path = Path(directory) / "legacy-map.xlsx"
            with patch("scrape_zj_land.fetch_resource_detail", return_value={"assignmentArea": 20000}), patch(
                "scrape_zj_land.normalize_resource_coordinate",
                return_value={"center": {"lng": 121.0, "lat": 30.0}, "locationType": "fixture"},
            ), patch("scrape_zj_land.load_map_config", return_value={"amapEnabled": False}), patch(
                "scrape_zj_land.browser_script", return_value=""
            ), patch("builtins.print"):
                save_to_excel([record], output_path, include_coords=True, include_map=True)

            points = json.loads(Path(directory, "legacy-map_coords.json").read_text(encoding="utf-8"))
            item = points[0]["items"][0]
            self.assertEqual(item["dealTotalPriceWan"], 219.35)
            self.assertEqual(item["dealUnitPriceYuanSqm"], 109.68)
            self.assertEqual(item["dealPriceMetadata"]["rawUnit"], "万元")
            map_html = Path(directory, "legacy-map_map.html").read_text(encoding="utf-8")
            self.assertIn("成交价来源与换算", map_html)
            self.assertIn("priceEvidence", map_html)

    def test_legacy_area_matches_modern_priority_and_preserves_negative_values(self):
        detail = {
            "assignmentArea": 10000,
            "dealPrice": 100,
            "dealPriceList": [{"price": 100, "unit": "万元"}],
        }
        record = {
            "_sourceEndpoint": "landbidding", "sourceId": "area-priority",
            "sourceCode": "area-priority", "landArea": 20000, "cjj": 100, "cjfs": "万元",
        }
        modern = enrich_record(record, lambda _source_id: detail)
        with tempfile.TemporaryDirectory() as directory:
            output_path = Path(directory) / "area-priority.xlsx"
            with patch("scrape_zj_land.fetch_resource_detail", return_value=detail), patch("builtins.print"):
                save_to_excel([record], output_path, include_coords=True)
            workbook = openpyxl.load_workbook(output_path, read_only=True, data_only=True)
            values = list(workbook.active.values)
            workbook.close()
            row = dict(zip(values[0], values[1]))
            coords = json.loads(Path(directory, "area-priority_coords.json").read_text(encoding="utf-8"))

        self.assertEqual(modern["area_sqm"], 20000.0)
        self.assertEqual(modern["deal_price"], 50.0)
        self.assertEqual(row["土地面积(平方米)"], modern["area_sqm"])
        self.assertEqual(row["成交单价(元/平方米)"], modern["deal_price"])
        self.assertEqual(coords[0]["_assignment_area_sqm"], modern["area_sqm"])

        zero_record = {
            "_sourceEndpoint": "landbidding", "sourceId": "zero-area",
            "sourceCode": "zero-area", "landArea": 0, "cjj": 100, "cjfs": "万元",
        }
        zero_detail = {
            "assignmentArea": 10000,
            "dealPrice": 100,
            "dealPriceList": [{"price": 100, "unit": "万元"}],
        }
        zero_modern = enrich_record(zero_record, lambda _source_id: zero_detail)
        with tempfile.TemporaryDirectory() as directory:
            output_path = Path(directory) / "zero-area.xlsx"
            with patch("scrape_zj_land.fetch_resource_detail", return_value=zero_detail), patch("builtins.print"):
                save_to_excel([zero_record], output_path, include_coords=True)
            workbook = openpyxl.load_workbook(output_path, read_only=True, data_only=True)
            values = list(workbook.active.values)
            workbook.close()
            row = dict(zip(values[0], values[1]))
            coords = json.loads(Path(directory, "zero-area_coords.json").read_text(encoding="utf-8"))

        self.assertEqual(zero_modern["area_sqm"], 0.0)
        self.assertIn(zero_modern["deal_price"], (None, ""))
        self.assertEqual(zero_modern["deal_total_price"], 100.0)
        self.assertEqual(row["土地面积(平方米)"], 0.0)
        self.assertIn(row["成交单价(元/平方米)"], (None, ""))
        self.assertEqual(row["成交总价(万元)"], 100.0)
        self.assertEqual(coords[0]["_assignment_area_sqm"], 0.0)

        negative_record = {
            "_sourceEndpoint": "landbidding", "sourceId": "negative-area",
            "sourceCode": "negative-area", "landArea": "-20,000", "cjj": 100, "cjfs": "万元",
        }
        negative_detail = {
            "assignmentArea": "-20000",
            "dealPrice": 100,
            "dealPriceList": [{"price": 100, "unit": "万元"}],
        }
        negative_modern = enrich_record(negative_record, lambda _source_id: negative_detail)
        with tempfile.TemporaryDirectory() as directory:
            output_path = Path(directory) / "negative-area.xlsx"
            with patch("scrape_zj_land.fetch_resource_detail", return_value=negative_detail), patch("builtins.print"):
                save_to_excel([negative_record], output_path, include_coords=True)
            workbook = openpyxl.load_workbook(output_path, read_only=True, data_only=True)
            values = list(workbook.active.values)
            workbook.close()
            row = dict(zip(values[0], values[1]))
            coords = json.loads(Path(directory, "negative-area_coords.json").read_text(encoding="utf-8"))

        self.assertEqual(negative_modern["area_sqm"], -20000.0)
        self.assertIn(negative_modern["deal_price"], (None, ""))
        self.assertEqual(negative_modern["deal_total_price"], 100.0)
        self.assertEqual(row["土地面积(平方米)"], -20000.0)
        self.assertIn(row["成交单价(元/平方米)"], (None, ""))
        self.assertEqual(row["成交总价(万元)"], 100.0)
        self.assertIn("负值", row["成交价警告"])
        self.assertEqual(coords[0]["_assignment_area_sqm"], -20000.0)

    def test_modern_excel_marks_formula_like_price_metadata_as_text(self):
        enriched = enrich_record({
            "startPrice": "=HYPERLINK(\"https://example.invalid\",\"open\")",
            "startPriceUnit": "万元",
            "landArea": 1000,
        }, lambda _source_id: {})
        row = row_from_item(enriched)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "modern-export.xlsx"
            write_excel(path, [row], {"maxPages": 1}, {})
            workbook = openpyxl.load_workbook(path, data_only=False)
            sheet = workbook.active
            column = [cell.value for cell in sheet[1]].index("起始价原始金额") + 1
            self.assertEqual(sheet.cell(2, column).value, "=HYPERLINK(\"https://example.invalid\",\"open\")")
            self.assertEqual(sheet.cell(2, column).data_type, "s")
            workbook.close()


if __name__ == "__main__":
    unittest.main()
