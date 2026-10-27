#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Native Messaging workbench adapter for the equipment class/life skill.

The skill intentionally leaves classification decisions to a human reviewer.
This adapter provides workbook paging, local reference lookup, safe-copy output,
the skill's consistency validator, and the seven-section Word report.
"""
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import uuid
import zipfile
from datetime import datetime
from pathlib import Path

SCRIPT_DIR = str(Path(__file__).resolve().parent)
if SCRIPT_DIR not in sys.path:
    sys.path.insert(0, SCRIPT_DIR)

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

DEFAULT_SHEET = "AI设备分类与年限"
REFERENCE_FILES = (
    "values.md",
    "table-2012.md",
    "table-2024.md",
    "2024-tables-raw.txt",
    "installation-rates.md",
)
ALLOWED_LARGE_CATEGORIES = (
    "家具",
    "文教、工美、体育和娱乐用品",
    "塑料制品",
    "通用设备",
    "专用设备",
    "电气机械及器材",
    "通信设备、计算机及其他电子设备",
    "仪器仪表",
)
SOURCE_LABELS = {
    "A": "《常用评估参数2012v1（初稿）》设备参考年限表",
    "B": "《资产评估报告编写模块常用语参考（2024版）》机器设备经济寿命参考年限表",
    "C": "网络查询：《通用与专用设备经济寿命参考年限表》",
    "D": "网络查询及行业惯例",
    "E": "网络查询：《GB/T 7723-2017》电子汽车衡耐用年限10-15年",
}


def emit(value):
    print("ATB_CLASS_LIFE_JSON:" + json.dumps(value, ensure_ascii=False, separators=(",", ":")))


def fail(reason, message):
    emit({"ok": False, "reason": reason, "message": message})
    return 1


def require_openpyxl():
    try:
        import openpyxl  # noqa: F401
    except Exception as error:
        raise RuntimeError("CLASS_LIFE_DEPENDENCY_MISSING:openpyxl") from error


def normalize_path(value, must_exist=True):
    raw = str(value or "").strip()
    if not raw or "\0" in raw or not os.path.isabs(raw):
        raise ValueError("CLASS_LIFE_PATH_INVALID")
    resolved = os.path.realpath(raw)
    if must_exist and not os.path.exists(resolved):
        raise FileNotFoundError("CLASS_LIFE_PATH_NOT_FOUND")
    return resolved


def workbook_sheet(workbook, requested=DEFAULT_SHEET):
    if requested in workbook.sheetnames:
        return workbook[requested]
    for sheet in workbook.worksheets:
        headers = [str(sheet.cell(1, col).value or "").strip() for col in range(1, 8)]
        if "设备名称" in headers[:2] and any("经济耐用年限" in value for value in headers):
            return sheet
    raise ValueError("CLASS_LIFE_TARGET_SHEET_NOT_FOUND")


def formula_and_link_status(path):
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=False)
    formula_count = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formula_count += 1
    external_links = len(getattr(wb, "_external_links", []) or [])
    names = list(getattr(wb, "defined_names", {}).keys())
    wb.close()
    with zipfile.ZipFile(path) as archive:
        members = archive.namelist()
        media_count = sum(name.startswith("xl/media/") for name in members)
        has_vba = any(name.endswith("vbaProject.bin") for name in members)
    return {
        "formulaCount": formula_count,
        "externalLinkCount": external_links,
        "definedNameCount": len(names),
        "mediaCount": media_count,
        "hasVba": has_vba,
        "requiresNativeEditor": bool(formula_count or external_links or has_vba),
    }


def read_inventory(path, offset, limit, sheet_name=DEFAULT_SHEET):
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=False)
    ws = workbook_sheet(wb, sheet_name)
    headers = [ws.cell(1, col).value for col in range(1, 8)]
    rows = []
    total = 0
    for row_number in range(2, ws.max_row + 1):
        values = [ws.cell(row_number, col).value for col in range(1, 8)]
        if values[0] is None and values[1] is None:
            continue
        if total >= offset and len(rows) < limit:
            rows.append({
                "rowNumber": row_number,
                "sequence": values[0],
                "name": str(values[1] or ""),
                "classification": values[2],
                "largeCategory": values[3],
                "economicLife": values[4],
                "lifeRange": values[5],
                "source": values[6],
            })
        total += 1
    title = ws.title
    wb.close()
    return {"sheetName": title, "headers": headers, "totalRows": total, "offset": offset, "limit": limit, "rows": rows}


def references_root():
    return Path(__file__).resolve().parent.parent / "references"


def lookup_reference(query, max_results=16):
    query = str(query or "").strip()
    if len(query) < 2:
        return {"query": query, "results": [], "message": "请输入至少两个字符进行本地口径检索。"}
    results = []
    normalized = re.sub(r"\s+", "", query).casefold()
    for file_name in REFERENCE_FILES:
        source = references_root() / file_name
        if not source.is_file():
            continue
        try:
            lines = source.read_text(encoding="utf-8", errors="replace").splitlines()
        except OSError:
            continue
        for line_number, line in enumerate(lines, 1):
            haystack = re.sub(r"\s+", "", line).casefold()
            if normalized in haystack:
                results.append({"file": file_name, "line": line_number, "text": line.strip()[:600]})
                if len(results) >= max_results:
                    return {"query": query, "results": results, "message": "仅返回前几条本地参考命中；仍需人工核对来源与设备口径。"}
    return {
        "query": query,
        "results": results,
        "message": "本地参考仅用于查找，不是自动分类结论；没有命中时请人工核查并记录真实依据。" if results else "本地参考未命中。请勿按相似名称自动套值，需人工查证后录入。",
    }


def session_directory(session_id):
    if not re.fullmatch(r"[0-9a-f]{32}", str(session_id or "")):
        raise ValueError("CLASS_LIFE_SESSION_INVALID")
    root = Path(tempfile.gettempdir()).resolve() / "AppraisalToolbox" / "equipment-class-life"
    target = (root / session_id).resolve()
    if not str(target).startswith(str(root) + os.sep):
        raise ValueError("CLASS_LIFE_SESSION_INVALID")
    return target


def load_session(session_id):
    directory = session_directory(session_id)
    meta_path = directory / "session.json"
    if not meta_path.is_file():
        raise FileNotFoundError("CLASS_LIFE_SESSION_EXPIRED")
    meta = json.loads(meta_path.read_text(encoding="utf-8"))
    if time.time() - float(meta.get("updatedAtEpoch", 0)) > 86400:
        shutil.rmtree(directory, ignore_errors=True)
        raise FileNotFoundError("CLASS_LIFE_SESSION_EXPIRED")
    return directory, meta


def save_session(directory, meta):
    meta["updatedAtEpoch"] = time.time()
    meta["updatedAt"] = datetime.now().astimezone().isoformat(timespec="seconds")
    path = directory / "session.json"
    temporary = path.with_suffix(".tmp")
    temporary.write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")
    os.replace(temporary, path)


def load_edits(directory):
    path = directory / "edits.json"
    if not path.is_file():
        return {}
    value = json.loads(path.read_text(encoding="utf-8"))
    return value if isinstance(value, dict) else {}


def store_page_edits(directory, rows):
    if not isinstance(rows, list) or len(rows) > 100:
        raise ValueError("CLASS_LIFE_EDIT_BATCH_INVALID")
    edits = load_edits(directory)
    allowed = {"classification", "largeCategory", "economicLife", "lifeRange", "source", "note"}
    for item in rows:
        if not isinstance(item, dict):
            continue
        row_number = int(item.get("rowNumber", 0))
        if row_number < 2 or row_number > 1_000_000:
            raise ValueError("CLASS_LIFE_ROW_INVALID")
        value = {key: item.get(key) for key in allowed if key in item}
        edits[str(row_number)] = value
    path = directory / "edits.json"
    temporary = path.with_suffix(".tmp")
    temporary.write_text(json.dumps(edits, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    os.replace(temporary, path)
    return len(edits)


def create_session(workbook_path, sheet_name):
    require_openpyxl()
    path = normalize_path(workbook_path)
    if Path(path).suffix.lower() not in (".xlsx", ".xlsm"):
        raise ValueError("CLASS_LIFE_XLSX_REQUIRED")
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=False)
    ws = workbook_sheet(wb, sheet_name)
    actual_sheet = ws.title
    wb.close()
    safety = formula_and_link_status(path)
    session_id = uuid.uuid4().hex
    directory = session_directory(session_id)
    directory.mkdir(parents=True, exist_ok=False, mode=0o700)
    meta = {
        "schemaVersion": 1,
        "sessionId": session_id,
        "workbookPath": path,
        "sheetName": actual_sheet,
        "safety": safety,
        "createdAt": datetime.now().astimezone().isoformat(timespec="seconds"),
        "createdAtEpoch": time.time(),
    }
    save_session(directory, meta)
    page = read_inventory(path, 0, 25, actual_sheet)
    return {"ok": True, "sessionId": session_id, "safety": safety, **page}


def apply_with_com(target_path, sheet_name, edits):
    try:
        import pythoncom
        import win32com.client as win32
    except Exception as error:
        raise RuntimeError("CLASS_LIFE_NATIVE_EDITOR_REQUIRED") from error
    pythoncom.CoInitialize()
    app = None
    workbook = None
    editor_name = ""
    try:
        for prog_id, label in (("KET.Application", "WPS"), ("Excel.Application", "Excel")):
            try:
                app = win32.DispatchEx(prog_id)
                editor_name = label
                break
            except Exception:
                app = None
        if app is None:
            raise RuntimeError("CLASS_LIFE_NATIVE_EDITOR_REQUIRED")
        app.Visible = False
        app.DisplayAlerts = False
        workbook = app.Workbooks.Open(os.path.abspath(target_path), UpdateLinks=0, ReadOnly=False, AddToMru=False)
        sheet = workbook.Worksheets(sheet_name)
        for row_key, values in edits.items():
            row_number = int(row_key)
            for column, key in zip((3, 4, 5, 6, 7), ("classification", "largeCategory", "economicLife", "lifeRange", "source")):
                if key not in values:
                    continue
                value = values.get(key)
                if column == 5 and value not in (None, ""):
                    value = int(float(value))
                if key in ("classification", "largeCategory", "lifeRange", "source") and value not in (None, ""):
                    # WPS/Excel 会把 "8-12" 这类区间文本自动转成日期；先强制文本
                    # 格式，避免 COM 写入损坏年限区间。
                    sheet.Cells(row_number, column).NumberFormat = "@"
                sheet.Cells(row_number, column).Value = value
        workbook.Save()
    except Exception:
        if workbook is None:
            raise
        raise RuntimeError("CLASS_LIFE_NATIVE_EDITOR_FAILED")
    finally:
        try:
            if workbook is not None:
                workbook.Close(SaveChanges=False)
        except Exception:
            pass
        try:
            if app is not None:
                app.Quit()
        except Exception:
            pass
        pythoncom.CoUninitialize()
    return editor_name


def apply_edits(source_path, target_path, sheet_name, edits, safety):
    from openpyxl import load_workbook

    shutil.copy2(source_path, target_path)
    if safety.get("requiresNativeEditor"):
        try:
            editor_name = apply_with_com(target_path, sheet_name, edits)
        except Exception:
            try:
                os.remove(target_path)
            except OSError:
                pass
            raise
        return f"native-{editor_name.lower()}"
    wb = load_workbook(target_path, keep_vba=Path(target_path).suffix.lower() == ".xlsm")
    ws = workbook_sheet(wb, sheet_name)
    for row_key, values in edits.items():
        row_number = int(row_key)
        for column, key in zip((3, 4, 5, 6, 7), ("classification", "largeCategory", "economicLife", "lifeRange", "source")):
            if key not in values:
                continue
            value = values.get(key)
            if column == 5 and value not in (None, ""):
                value = int(float(value))
            ws.cell(row_number, column).value = value
    wb.save(target_path)
    wb.close()
    return "openpyxl-copy"


def validate_workbook(path, sheet_name):
    verifier = Path(__file__).resolve().parent / "verify_fill.py"
    result = subprocess.run(
        [sys.executable, "-X", "utf8", "-u", str(verifier), str(path), "--sheet", sheet_name],
        capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=120,
    )
    output = (result.stdout or "").strip()
    errors = [line.strip().lstrip("- ") for line in output.splitlines() if line.strip().startswith(" - ")]
    return {"ok": result.returncode == 0, "exitCode": result.returncode, "output": output[-18000:], "issues": errors[:300]}


def staging_copy(directory, source_path, sheet_name, edits, safety):
    extension = Path(source_path).suffix.lower()
    path = directory / f"validation-copy{extension if extension in ('.xlsx', '.xlsm') else '.xlsx'}"
    if path.exists():
        path.unlink()
    apply_edits(source_path, str(path), sheet_name, edits, safety)
    return path


def report_rows(workbook_path, sheet_name):
    from openpyxl import load_workbook

    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    ws = workbook_sheet(wb, sheet_name)
    rows = []
    for r in range(2, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, 8)]
        if vals[0] is None and vals[1] is None:
            continue
        rows.append({
            "sequence": vals[0], "name": str(vals[1] or ""), "classification": vals[2],
            "largeCategory": vals[3], "economicLife": vals[4], "lifeRange": vals[5], "source": vals[6],
        })
    wb.close()
    return rows


def create_report(output_path, rows, edits, validation, source_workbook, sheet_name):
    try:
        from docx import Document
        from docx.shared import Pt
    except Exception as error:
        raise RuntimeError("CLASS_LIFE_DEPENDENCY_MISSING:python-docx") from error
    document = Document()
    document.add_heading("设备分类与年限补充说明及一致性校验报告", 0)
    document.add_paragraph("工作簿：" + os.path.basename(source_workbook))
    document.add_paragraph("工作表：" + sheet_name)
    document.add_paragraph("生成时间：" + datetime.now().astimezone().strftime("%Y-%m-%d %H:%M:%S %z"))
    document.add_heading("一、任务完成概况", level=1)
    document.add_paragraph(f"设备行数：{len(rows):,}；一致性校验：{'通过' if validation['ok'] else '未通过'}。")
    document.add_paragraph("本模块不自动作出无法由本地资料确认的分类结论；分类和年限由操作人员复核录入。")
    distribution = {}
    for item in rows:
        key = str(item.get("largeCategory") or "未填写")
        distribution[key] = distribution.get(key, 0) + 1
    document.add_paragraph("指数-设备大类分布：" + "；".join(f"{key} {value:,} 台" for key, value in sorted(distribution.items())))

    document.add_heading("二、分类与年限口径汇总", level=1)
    groups = {}
    for item in rows:
        key = (str(item.get("classification") or "未填写"), str(item.get("largeCategory") or "未填写"),
               str(item.get("economicLife") or "未填写"), str(item.get("lifeRange") or "未填写"), str(item.get("source") or "未填写"))
        groups[key] = groups.get(key, 0) + 1
    table = document.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, ("分类", "指数-设备大类", "年限", "区间", "来源依据", "台数")):
        cell.text = value
    for key, count in sorted(groups.items()):
        for cell, value in zip(table.add_row().cells, (*key, f"{count:,}")):
            cell.text = value

    document.add_heading("三、异名同物识别与合并口径", level=1)
    from verify_fill import base, group_key
    synonym_groups = {}
    for item in rows:
        name = base(item["name"])
        synonym_groups.setdefault(group_key(name), []).append(name)
    merged = [(key, values) for key, values in synonym_groups.items() if len(set(values)) > 1]
    document.add_paragraph("按技能内同义词与基名规则识别到的合并组：" + (str(len(merged)) if merged else "0") + " 组。")
    for key, values in merged[:120]:
        document.add_paragraph(f"{key}：" + "、".join(sorted(set(values))), style="List Bullet")

    document.add_heading("四、同种设备一致性校验", level=1)
    document.add_paragraph(validation["output"] or "未生成校验输出。")

    document.add_heading("五、主要依据", level=1)
    used_sources = sorted({str(item.get("source") or "").strip() for item in rows if str(item.get("source") or "").strip()})
    for source in used_sources:
        document.add_paragraph(source, style="List Bullet")
    document.add_paragraph("分类-年限口径与异名同物参考：技能包 references/values.md、table-2012.md、table-2024.md。")

    document.add_heading("六、特殊事项与复核提示", level=1)
    saved_edits = edits
    for row_key, value in sorted(saved_edits.items(), key=lambda pair: int(pair[0])):
        note = str(value.get("note") or "").strip()
        if note:
            document.add_paragraph(f"行 {row_key}：{note}", style="List Bullet")
    if not any(str(value.get("note") or "").strip() for value in saved_edits.values()):
        document.add_paragraph("未录入需要专项裁量的备注。请结合技能提示复核不动产/场站、劳动防护用品、油烟净化器、电动门及空调等特殊项。")

    document.add_heading("七、跨批次口径比对", level=1)
    document.add_paragraph("本次未提供已认可历史输出 TSV，跨批次比对未执行。")
    document.save(output_path)


def operation(request):
    op = str(request.get("operation") or "").strip()
    if op == "lookup":
        return {"ok": True, **lookup_reference(request.get("query"))}

    require_openpyxl()

    if op == "create":
        return create_session(request.get("workbookPath"), request.get("sheetName") or DEFAULT_SHEET)

    session_id = request.get("sessionId")
    directory, meta = load_session(session_id)
    if op == "page":
        page = read_inventory(meta["workbookPath"], max(0, int(request.get("offset", 0))),
                              min(100, max(1, int(request.get("limit", 25)))), meta["sheetName"])
        edits = load_edits(directory)
        editable_fields = ("classification", "largeCategory", "economicLife", "lifeRange", "source", "note")
        for row in page["rows"]:
            saved = edits.get(str(row["rowNumber"])) or {}
            if isinstance(saved, dict):
                for field in editable_fields:
                    if field in saved:
                        row[field] = saved[field]
        return {"ok": True, "sessionId": session_id, "safety": meta["safety"], **page}
    if op == "save_page":
        count = store_page_edits(directory, request.get("rows") or [])
        save_session(directory, meta)
        return {"ok": True, "savedRows": count, "sessionId": session_id}

    edits = load_edits(directory)
    source_path = meta["workbookPath"]
    sheet_name = meta["sheetName"]
    if op == "validate":
        target = staging_copy(directory, source_path, sheet_name, edits, meta["safety"])
        try:
            result = validate_workbook(str(target), sheet_name)
        finally:
            try:
                target.unlink()
            except OSError:
                pass
        save_session(directory, meta)
        return {"ok": result["ok"], "validation": result, "changedRows": len(edits), "sessionId": session_id}

    if op == "export":
        output_directory = normalize_path(request.get("outputDirectory"))
        if not os.path.isdir(output_directory):
            raise ValueError("CLASS_LIFE_OUTPUT_DIRECTORY_INVALID")
        stamp = datetime.now().strftime("%Y%m%d-%H%M%S-%f")
        stem = Path(source_path).stem
        extension = Path(source_path).suffix.lower()
        output_extension = extension if extension in (".xlsx", ".xlsm") else ".xlsx"
        output_workbook = os.path.join(output_directory, f"{stem}_设备分类与年限_{stamp}{output_extension}")
        if os.path.exists(output_workbook):
            raise FileExistsError("CLASS_LIFE_OUTPUT_EXISTS")
        mode = apply_edits(source_path, output_workbook, sheet_name, edits, meta["safety"])
        validation = validate_workbook(output_workbook, sheet_name)
        if not validation["ok"]:
            try:
                os.remove(output_workbook)
            except OSError:
                pass
            return {"ok": False, "reason": "CLASS_LIFE_VALIDATION_FAILED", "validation": validation,
                    "message": "一致性校验未通过，未保留最终成果文件。请按问题清单修正后重试。"}
        rows = report_rows(output_workbook, sheet_name)
        report_path = os.path.join(output_directory, f"{stem}_设备分类与年限报告_{stamp}.docx")
        try:
            create_report(report_path, rows, edits, validation, source_path, sheet_name)
        except Exception:
            try:
                os.remove(output_workbook)
            except OSError:
                pass
            try:
                os.remove(report_path)
            except OSError:
                pass
            raise
        save_session(directory, meta)
        return {"ok": True, "workbookPath": output_workbook, "reportPath": report_path,
                "changedRows": len(edits), "validation": validation, "writeMode": mode,
                "message": "已生成新工作簿副本和 Word 复核报告；原始文件未改动。"}

    if op == "close":
        shutil.rmtree(directory, ignore_errors=True)
        return {"ok": True, "closed": True, "message": "本次暂存数据已清除。"}
    raise ValueError("CLASS_LIFE_OPERATION_INVALID")


def main():
    args = sys.argv[1:]
    if "--request-file" not in args:
        return fail("CLASS_LIFE_REQUEST_INVALID", "缺少模块请求。")
    try:
        request_path = args[args.index("--request-file") + 1]
        request_path = normalize_path(request_path)
        request = json.loads(Path(request_path).read_text(encoding="utf-8"))
        result = operation(request)
        emit(result)
        return 0 if result.get("ok") else 1
    except Exception as error:
        message = str(error)
        reason = message.split(":", 1)[0] if message else "CLASS_LIFE_OPERATION_FAILED"
        friendly = {
            "CLASS_LIFE_DEPENDENCY_MISSING": "本机缺少设备分类模块所需的 Python 组件。",
            "CLASS_LIFE_TARGET_SHEET_NOT_FOUND": "未找到“AI设备分类与年限”工作表或兼容的七列表头。",
            "CLASS_LIFE_SESSION_EXPIRED": "本次设备清单暂存已过期，请重新选择工作簿。",
            "CLASS_LIFE_NATIVE_EDITOR_REQUIRED": "工作簿含公式、外部链接或宏；当前环境没有可用的 WPS/Excel 原生编辑组件，未保存结果。",
            "CLASS_LIFE_NATIVE_EDITOR_FAILED": "WPS/Excel 原生编辑失败，未保留部分输出文件。",
            "CLASS_LIFE_PATH_INVALID": "选择的文件或目录路径无效。",
            "CLASS_LIFE_PATH_NOT_FOUND": "选择的文件或目录已不存在。",
            "CLASS_LIFE_XLSX_REQUIRED": "请选择 .xlsx 或 .xlsm 设备工作簿。",
            "CLASS_LIFE_VALIDATION_FAILED": "一致性校验未通过，未生成最终成果。",
        }
        return fail(reason, friendly.get(reason, "设备分类模块操作失败，请检查工作簿和输出目录后重试。"))


if __name__ == "__main__":
    sys.exit(main())
